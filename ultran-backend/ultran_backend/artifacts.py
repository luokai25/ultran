"""Workspace-scoped file artifact generation and storage."""
from __future__ import annotations

import csv
import hashlib
import html
import io
import json
import re
import uuid
import zipfile
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path
import threading
from typing import Any
from xml.sax.saxutils import escape

from .errors import NotFound, ValidationFailed

FORMATS = {
    "html": ("text/html; charset=utf-8", "HTML document"),
    "md": ("text/markdown; charset=utf-8", "Markdown document"),
    "txt": ("text/plain; charset=utf-8", "Text document"),
    "json": ("application/json", "JSON data"),
    "csv": ("text/csv; charset=utf-8", "CSV spreadsheet"),
    "svg": ("image/svg+xml", "SVG image"),
    "pdf": ("application/pdf", "PDF document"),
    "pptx": ("application/vnd.openxmlformats-officedocument.presentationml.presentation", "PowerPoint presentation"),
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "Word document"),
    "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "Excel workbook"),
}


def _safe_name(value: str) -> str:
    value = re.sub(r"[^a-zA-Z0-9._-]+", "-", value.strip()).strip(".-_")[:64]
    return value or "ultran-artifact"


class _SafeHTML(HTMLParser):
    ALLOWED = {"main", "header", "footer", "nav", "section", "article", "div", "span", "h1", "h2", "h3", "h4", "h5", "h6", "p", "br", "hr", "strong", "b", "em", "i", "u", "small", "blockquote", "pre", "code", "ul", "ol", "li", "table", "thead", "tbody", "tr", "th", "td", "a", "figure", "figcaption", "details", "summary"}
    VOID = {"br", "hr"}
    DROP_CONTENT = {"script", "style", "iframe", "object", "embed", "form", "input", "button", "textarea", "video", "audio"}
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.drop_depth = 0
        self.stack: list[str] = []
    def handle_starttag(self, tag, attrs):
        if tag in self.DROP_CONTENT:
            self.drop_depth += 1
            return
        if self.drop_depth or tag not in self.ALLOWED:
            return
        safe_attrs = []
        for key, value in attrs:
            if value is None or key not in {"href", "title", "alt", "class", "id", "colspan", "rowspan", "lang"}:
                continue
            if key == "href" and not (value.startswith(("https://", "http://", "mailto:", "#", "/")) and not value.startswith("//")):
                continue
            safe_attrs.append(f' {key}="{html.escape(value, quote=True)}"')
        self.parts.append("<" + tag + "".join(safe_attrs) + ">")
        if tag not in self.VOID:
            self.stack.append(tag)
    def handle_endtag(self, tag):
        if tag in self.DROP_CONTENT and self.drop_depth:
            self.drop_depth -= 1
            return
        if self.drop_depth or tag not in self.ALLOWED or tag in self.VOID:
            return
        if tag in self.stack:
            while self.stack:
                current = self.stack.pop()
                self.parts.append(f"</{current}>")
                if current == tag:
                    break
    def handle_data(self, data):
        if not self.drop_depth:
            self.parts.append(html.escape(data))


def _html_bytes(title: str, content: str) -> bytes:
    source = content.strip()
    if source.startswith("```"):
        source = re.sub(r"^```(?:html)?\s*|\s*```$", "", source, flags=re.I)
    if "<" in source and ">" in source:
        parser = _SafeHTML()
        parser.feed(source)
        body = "".join(parser.parts)
        while parser.stack:
            body += f"</{parser.stack.pop()}>"
    else:
        blocks = []
        for paragraph in source.split("\n\n"):
            lines = paragraph.splitlines()
            if lines and lines[0].startswith("#"):
                level = min(6, max(1, len(lines[0]) - len(lines[0].lstrip("#"))))
                blocks.append(f"<h{level}>{html.escape(lines[0].lstrip('# ').strip())}</h{level}>")
                if len(lines) > 1:
                    blocks.append("<p>" + html.escape("\n".join(lines[1:])).replace("\n", "<br>") + "</p>")
            else:
                blocks.append("<p>" + html.escape(paragraph).replace("\n", "<br>") + "</p>")
        body = "".join(blocks)
    document = ("<!doctype html><html lang='en'><head><meta charset='utf-8'>"
                "<meta name='viewport' content='width=device-width,initial-scale=1'><title>" + html.escape(title) +
                "</title><style>body{margin:0;background:#101014;color:#eee;font:16px/1.7 system-ui}main{max-width:920px;margin:40px auto;padding:32px}h1,h2,h3{color:#e3c15c}a{color:#e3c15c}table{border-collapse:collapse}td,th{padding:8px;border:1px solid #555}pre{white-space:pre-wrap;background:#222;padding:14px;border-radius:8px}</style></head><body><main>" + body + "</main></body></html>")
    return document.encode("utf-8")


def _library_pdf_bytes(title: str, content: str) -> bytes | None:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
    except ImportError:
        return None
    regular_font = "Helvetica"
    bold_font = "Helvetica-Bold"
    font_candidates = [
        (Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")),
        (Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"), Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf")),
    ]
    for regular_path, bold_path in font_candidates:
        if regular_path.is_file():
            pdfmetrics.registerFont(TTFont("UltranSans", str(regular_path)))
            regular_font = "UltranSans"
            if bold_path.is_file():
                pdfmetrics.registerFont(TTFont("UltranSansBold", str(bold_path)))
                bold_font = "UltranSansBold"
            else:
                bold_font = regular_font
            break
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, rightMargin=20 * mm, leftMargin=20 * mm,
                            topMargin=20 * mm, bottomMargin=20 * mm, title=title)
    base = getSampleStyleSheet()
    title_style = ParagraphStyle("UltranTitle", parent=base["Title"], fontName=bold_font, textColor=colors.HexColor("#9B761B"), alignment=TA_LEFT, spaceAfter=14)
    heading = ParagraphStyle("UltranHeading", parent=base["Heading2"], fontName=bold_font, textColor=colors.HexColor("#765B16"), spaceBefore=12, spaceAfter=6)
    body = ParagraphStyle("UltranBody", parent=base["BodyText"], fontName=regular_font, leading=15, spaceAfter=7)
    story = [Paragraph(escape(title), title_style)]
    for line in content.splitlines():
        stripped = line.strip()
        if not stripped:
            story.append(Spacer(1, 4))
        elif stripped.startswith("### "):
            story.append(Paragraph(escape(stripped[4:]), heading))
        elif stripped.startswith("## ") or stripped.startswith("# "):
            story.append(Paragraph(escape(stripped.lstrip("# ")), heading))
        elif stripped.startswith(("- ", "* ")):
            story.append(Paragraph("- " + escape(stripped[2:]), body))
        else:
            story.append(Paragraph(escape(stripped), body))
    doc.build(story)
    return output.getvalue()


def _library_pptx_bytes(title: str, content: str) -> bytes | None:
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
    except ImportError:
        return None
    sections: list[tuple[str, str]] = []
    heading, body = title, []
    for line in content.splitlines():
        if line.startswith("# ") and (body or not sections):
            if body or sections:
                sections.append((heading, "\n".join(body)))
            heading, body = line[2:].strip(), []
        else:
            body.append(line)
    if body or not sections:
        sections.append((heading, "\n".join(body)))
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    for slide_title, slide_body in sections[:30]:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = RGBColor(17, 17, 20)
        title_box = slide.shapes.add_textbox(Inches(.7), Inches(.45), Inches(12), Inches(1.0))
        title_frame = title_box.text_frame
        title_frame.text = slide_title or title
        title_run = title_frame.paragraphs[0].runs[0]
        title_run.font.name = "Aptos Display"
        title_run.font.size = Pt(30)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(212, 175, 55)
        body_box = slide.shapes.add_textbox(Inches(.85), Inches(1.7), Inches(11.6), Inches(5.1))
        frame = body_box.text_frame
        frame.word_wrap = True
        frame.clear()
        for index, line in enumerate(slide_body.splitlines()[:24]):
            if not line.strip():
                continue
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = line.lstrip("-*• ") if line.lstrip().startswith(("-", "*", "•")) else line
            paragraph.level = 0
            paragraph.font.name = "Aptos"
            paragraph.font.size = Pt(18)
            paragraph.font.color.rgb = RGBColor(238, 235, 224)
            if line.lstrip().startswith(("-", "*", "•")):
                paragraph.text = "• " + paragraph.text
        footer = slide.shapes.add_textbox(Inches(11.9), Inches(7.05), Inches(1), Inches(.25))
        footer.text_frame.text = f"{len(presentation.slides):02d}"
        footer.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
        footer.text_frame.paragraphs[0].font.size = Pt(9)
        footer.text_frame.paragraphs[0].font.color.rgb = RGBColor(150, 150, 150)
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def _library_docx_bytes(title: str, content: str) -> bytes | None:
    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
    except ImportError:
        return None
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(.75)
    section.bottom_margin = Inches(.75)
    document.add_heading(title, 0)
    for line in content.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("### "):
            document.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            document.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            document.add_heading(stripped[2:], level=1)
        elif stripped.startswith(("- ", "* ")):
            document.add_paragraph(stripped[2:], style="List Bullet")
        elif re.match(r"^\d+[.)]\s", stripped):
            document.add_paragraph(re.sub(r"^\d+[.)]\s", "", stripped), style="List Number")
        else:
            document.add_paragraph(stripped)
    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(10.5)
    document.styles["Title"].font.color.rgb = RGBColor(155, 118, 27)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _library_xlsx_bytes(content: str) -> bytes | None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        return None
    rows = list(csv.reader(io.StringIO(content))) or [[]]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ultran Data"
    for row_index, row in enumerate(rows[:10000], 1):
        for column_index, value in enumerate(row[:256], 1):
            cell = sheet.cell(row=row_index, column=column_index)
            text = value.strip()
            if row_index > 1 and re.fullmatch(r"-?(?:\d+\.?\d*|\.\d+)", text):
                cell.value = float(text) if "." in text else int(text)
            else:
                cell.value = value
                cell.data_type = "s"  # prevent spreadsheet formula injection
            if row_index == 1:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="9B761B")
                cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_index in range(1, min(sheet.max_column, 64) + 1):
        width = min(48, max(12, max((len(str(sheet.cell(r, column_index).value or "")) for r in range(1, min(sheet.max_row, 500) + 1)), default=10) + 2))
        sheet.column_dimensions[get_column_letter(column_index)].width = width
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _pdf_bytes(title: str, content: str) -> bytes:
    """Create a small, valid text PDF without requiring a system binary."""
    lines: list[str] = [title, ""]
    for paragraph in content.replace("\r", "").splitlines():
        while len(paragraph) > 88:
            cut = paragraph.rfind(" ", 0, 88)
            cut = cut if cut > 0 else 88
            lines.append(paragraph[:cut])
            paragraph = paragraph[cut:].lstrip()
        lines.append(paragraph)
    pages = [lines[i:i + 48] for i in range(0, max(1, len(lines)), 48)]
    objects: list[bytes] = []
    page_ids = []
    font_id = 3
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(b"")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for page in pages:
        page_id = len(objects) + 1
        content_id = page_id + 1
        page_ids.append(page_id)
        commands = ["BT", "/F1 10 Tf", "50 790 Td", "14 TL"]
        for line in page:
            safe = line.encode("latin-1", "replace").decode("latin-1")
            safe = safe.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            commands.append(f"({safe}) Tj T*" )
        commands.append("ET")
        stream = "\n".join(commands).encode("latin-1")
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>".encode())
        objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    objects[1] = ("<< /Type /Pages /Kids [" + " ".join(f"{i} 0 R" for i in page_ids) + f"] /Count {len(page_ids)} >>").encode()
    out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(out))
        out.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(out)
    out.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        out.extend(f"{offset:010d} 00000 n \n".encode())
    out.extend(f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return bytes(out)


def _pptx_bytes(title: str, content: str) -> bytes:
    """Create a minimal editable Open XML presentation with one slide per heading."""
    sections: list[tuple[str, str]] = []
    heading = title
    body: list[str] = []
    for line in content.splitlines():
        if line.startswith("# ") and (body or not sections):
            if body or sections:
                sections.append((heading, "\n".join(body)))
            heading, body = line[2:].strip(), []
        else:
            body.append(line)
    if body or not sections:
        sections.append((heading, "\n".join(body)))
    sections = sections[:30]
    def shape(text: str, x: int, y: int, w: int, h: int, sid: int, size: int, bold: bool = False) -> str:
        runs = "".join(f'<a:p><a:r><a:rPr lang="en-US" sz="{size}" b="{1 if bold else 0}"/><a:t>{escape(line)}</a:t></a:r></a:p>' for line in text.splitlines()[:28] if line.strip())
        return f'<p:sp><p:nvSpPr><p:cNvPr id="{sid}" name="Text {sid}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{w}" cy="{h}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/><a:ln><a:noFill/></a:ln></p:spPr><p:txBody><a:bodyPr wrap="square"/><a:lstStyle/>{runs}</p:txBody></p:sp>'
    files: dict[str, str] = {
        "[Content_Types].xml": '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/ppt/presentation.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/><Override PartName="/ppt/slideMasters/slideMaster1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slideMaster+xml"/><Override PartName="/ppt/slideLayouts/slideLayout1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slideLayout+xml"/><Override PartName="/ppt/theme/theme1.xml" ContentType="application/vnd.openxmlformats-officedocument.theme+xml"/>' + ''.join(f'<Override PartName="/ppt/slides/slide{i}.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>' for i in range(1, len(sections)+1)) + '</Types>',
        "_rels/.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="ppt/presentation.xml"/></Relationships>',
        "ppt/presentation.xml": '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><p:presentation xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:sldMasterIdLst><p:sldMasterId id="2147483648" r:id="rId1"/></p:sldMasterIdLst><p:sldIdLst>' + ''.join(f'<p:sldId id="{255+i}" r:id="rId{i+2}"/>' for i in range(len(sections))) + '</p:sldIdLst><p:sldSz cx="12192000" cy="6858000" type="wide"/><p:notesSz cx="6858000" cy="9144000"/></p:presentation>',
        "ppt/_rels/presentation.xml.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideMaster" Target="slideMasters/slideMaster1.xml"/>' + ''.join(f'<Relationship Id="rId{i+2}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide{i+1}.xml"/>' for i in range(len(sections))) + '</Relationships>',
        "ppt/slideMasters/slideMaster1.xml": '<p:sldMaster xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:cSld><p:spTree><p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr/></p:spTree></p:cSld><p:clrMap accent1="4472C4" accent2="ED7D31" accent3="A5A5A5" accent4="FFC000" accent5="5B9BD5" accent6="70AD47" bg1="lt1" tx1="dk1" bg2="lt2" tx2="dk2"/><p:txStyles><p:titleStyle/><p:bodyStyle/><p:otherStyle/></p:txStyles></p:sldMaster>',
        "ppt/slideLayouts/slideLayout1.xml": '<p:sldLayout xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" type="blank" preserve="1"><p:cSld name="Blank"><p:spTree><p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr/></p:spTree></p:cSld></p:sldLayout>',
        "ppt/slideMasters/_rels/slideMaster1.xml.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout" Target="../slideLayouts/slideLayout1.xml"/><Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme" Target="../theme/theme1.xml"/></Relationships>',
        "ppt/theme/theme1.xml": '<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" name="Ultran"><a:themeElements><a:clrScheme name="Ultran"><a:dk1><a:srgbClr val="111111"/></a:dk1><a:lt1><a:srgbClr val="FFFFFF"/></a:lt1><a:dk2><a:srgbClr val="222222"/></a:dk2><a:lt2><a:srgbClr val="F4F0E4"/></a:lt2><a:accent1><a:srgbClr val="D4AF37"/></a:accent1><a:accent2><a:srgbClr val="8A6D1F"/></a:accent2><a:accent3><a:srgbClr val="4472C4"/></a:accent3><a:accent4><a:srgbClr val="ED7D31"/></a:accent4><a:accent5><a:srgbClr val="70AD47"/></a:accent5><a:accent6><a:srgbClr val="5B9BD5"/></a:accent6><a:hlink><a:srgbClr val="0563C1"/></a:hlink><a:folHlink><a:srgbClr val="954F72"/></a:folHlink></a:clrScheme><a:fontScheme name="Ultran"><a:majorFont><a:latin typeface="Aptos Display"/><a:ea typeface=""/><a:cs typeface=""/></a:majorFont><a:minorFont><a:latin typeface="Aptos"/><a:ea typeface=""/><a:cs typeface=""/></a:minorFont></a:fontScheme><a:fmtScheme name="Ultran"><a:fillStyleLst><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:fillStyleLst><a:lnStyleLst><a:ln w="6350"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:ln></a:lnStyleLst><a:effectStyleLst><a:effectStyle><a:effectLst/></a:effectStyle></a:effectStyleLst><a:bgFillStyleLst><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:bgFillStyleLst></a:fmtScheme></a:themeElements><a:objectDefaults/><a:extraClrSchemeLst/></a:theme>',
        "ppt/slideLayouts/_rels/slideLayout1.xml.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideMaster" Target="../slideMasters/slideMaster1.xml"/></Relationships>',
    }
    for i, (heading, text) in enumerate(sections, 1):
        tree = '<p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr/>'
        tree += shape(heading, 500000, 300000, 11000000, 900000, 2, 2800, True)
        tree += shape(text, 650000, 1500000, 10800000, 4800000, 3, 1600)
        files[f"ppt/slides/slide{i}.xml"] = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:cSld><p:spTree>{tree}</p:spTree></p:cSld><p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr></p:sld>'
        files[f"ppt/slides/_rels/slide{i}.xml.rels"] = '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout" Target="../slideLayouts/slideLayout1.xml"/></Relationships>'
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for path, value in files.items():
            archive.writestr(path, value)
    return out.getvalue()


def _docx_bytes(title: str, content: str) -> bytes:
    paragraphs = [f"<w:p><w:r><w:t>{escape(line)}</w:t></w:r></w:p>" for line in [title, "", *content.splitlines()]]
    document = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>' + "".join(paragraphs) + '<w:sectPr><w:pgSz w:w="12240" w:h="15840"/><w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440"/></w:sectPr></w:body></w:document>'
    files = {
        "[Content_Types].xml": '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>',
        "_rels/.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>',
        "word/document.xml": document,
    }
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, value in files.items(): archive.writestr(name, value)
    return out.getvalue()


def _xlsx_bytes(content: str) -> bytes:
    rows = list(csv.reader(io.StringIO(content))) or [[]]
    def col_name(index: int) -> str:
        result = ""
        while index:
            index, rem = divmod(index - 1, 26)
            result = chr(65 + rem) + result
        return result
    cells = []
    for row_index, row in enumerate(rows[:10000], 1):
        parts = []
        for col_index, value in enumerate(row[:256], 1):
            ref = f"{col_name(col_index)}{row_index}"
            if re.fullmatch(r"-?(?:\d+\.?\d*|\.\d+)", value.strip()):
                parts.append(f'<c r="{ref}"><v>{value.strip()}</v></c>')
            else:
                parts.append(f'<c r="{ref}" t="inlineStr"><is><t>{escape(value)}</t></is></c>')
        cells.append(f'<row r="{row_index}">{"".join(parts)}</row>')
    sheet = '<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>' + "".join(cells) + '</sheetData></worksheet>'
    files = {
        "[Content_Types].xml": '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>',
        "_rels/.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>',
        "xl/workbook.xml": '<?xml version="1.0" encoding="UTF-8"?><workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Sheet1" sheetId="1" r:id="rId1"/></sheets></workbook>',
        "xl/_rels/workbook.xml.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/></Relationships>',
        "xl/worksheets/sheet1.xml": sheet,
    }
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, value in files.items(): archive.writestr(name, value)
    return out.getvalue()


class ArtifactService:
    def __init__(self, root: str | Path) -> None:
        self.root = Path(root).resolve()
        self.root.mkdir(parents=True, exist_ok=True)
        self._items: dict[str, dict[str, Any]] = {}
        self._lock = threading.RLock()
        for meta_path in self.root.glob("*/.*.meta.json"):
            try:
                item = json.loads(meta_path.read_text("utf-8"))
                file_path = Path(item["path"]).resolve()
                if file_path.is_file() and file_path.parent == meta_path.parent.resolve():
                    self._items[item["id"]] = item
            except (OSError, ValueError, KeyError, TypeError):
                continue

    def create(self, workspace_id: str, title: str, format: str, content: str) -> dict[str, Any]:
        fmt = format.lower().lstrip(".")
        if fmt not in FORMATS:
            raise ValidationFailed("unsupported artifact format", details={"format": fmt, "supported": sorted(FORMATS)})
        if not title.strip() or len(title) > 160:
            raise ValidationFailed("title must contain 1 to 160 characters")
        if len(content.encode("utf-8")) > 2 * 1024 * 1024:
            raise ValidationFailed("artifact content exceeds 2 MiB")
        artifact_id = f"art_{uuid.uuid4().hex[:16]}"
        workspace_key = hashlib.sha256(workspace_id.encode()).hexdigest()[:24]
        directory = (self.root / workspace_key).resolve()
        directory.mkdir(parents=True, exist_ok=True)
        filename = f"{_safe_name(title)}-{artifact_id[-6:]}.{fmt}"
        path = (directory / filename).resolve()
        if path.parent != directory:
            raise ValidationFailed("invalid artifact filename")
        if fmt == "html":
            payload = _html_bytes(title, content)
        elif fmt == "json":
            try:
                payload = json.dumps(json.loads(content), ensure_ascii=False, indent=2).encode()
            except ValueError as exc:
                raise ValidationFailed("JSON artifact content is not valid JSON") from exc
        elif fmt == "csv":
            rows = list(csv.reader(io.StringIO(content)))
            output = io.StringIO()
            csv.writer(output).writerows(rows)
            payload = output.getvalue().encode()
        elif fmt == "svg":
            if "<svg" not in content[:500].lower():
                raise ValidationFailed("SVG content must contain an svg root element")
            payload = content.encode()
        elif fmt == "pdf":
            payload = _library_pdf_bytes(title, content) or _pdf_bytes(title, content)
        elif fmt == "pptx":
            payload = _library_pptx_bytes(title, content) or _pptx_bytes(title, content)
        elif fmt == "docx":
            payload = _library_docx_bytes(title, content) or _docx_bytes(title, content)
        elif fmt == "xlsx":
            payload = _library_xlsx_bytes(content) or _xlsx_bytes(content)
        else:
            payload = content.encode("utf-8")
        path.write_bytes(payload)
        preview_path = ""
        if fmt == "pptx":
            preview_path = str(directory / f".{artifact_id}.preview.html")
            slide_sections = [part.strip() for part in re.split(r"(?m)^# ", content) if part.strip()]
            cards = []
            for index, section in enumerate(slide_sections[:30], 1):
                head, _, body = section.partition("\n")
                cards.append(f"<section class='slide'><small>SLIDE {index:02d}</small><h2>{html.escape(head or title)}</h2><p>{html.escape(body)}</p></section>")
            preview_doc = "<!doctype html><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><title>" + html.escape(title) + "</title><style>body{margin:0;padding:24px;background:#111;color:#eee;font:15px system-ui}.slide{max-width:900px;min-height:360px;margin:20px auto;padding:42px;background:#191919;border:1px solid #a88a36;border-radius:18px;box-shadow:0 12px 40px #000}.slide small{color:#d4af37;letter-spacing:.15em}.slide h2{font-size:30px}.slide p{white-space:pre-wrap;line-height:1.65}</style>" + "".join(cards)
            Path(preview_path).write_text(preview_doc, encoding="utf-8")
        elif fmt in {"docx", "xlsx", "csv", "md", "txt", "json"}:
            preview_path = str(directory / f".{artifact_id}.preview.html")
            if fmt in {"xlsx", "csv"}:
                rows = list(csv.reader(io.StringIO(content)))[:500]
                table_rows = []
                for row_index, row in enumerate(rows):
                    tag = "th" if row_index == 0 else "td"
                    table_rows.append("<tr>" + "".join(f"<{tag}>{html.escape(cell)}</{tag}>" for cell in row[:64]) + "</tr>")
                preview_body = "<table>" + "".join(table_rows) + "</table>"
            else:
                display = json.dumps(json.loads(content), ensure_ascii=False, indent=2) if fmt == "json" else content
                preview_body = "<pre>" + html.escape(display) + "</pre>"
            preview_doc = _html_bytes(title, "<h1>" + html.escape(title) + "</h1>" + preview_body).decode("utf-8")
            Path(preview_path).write_text(preview_doc, encoding="utf-8")
        metadata = {"id": artifact_id, "workspace_id": workspace_id, "title": title, "format": fmt,
                    "mime_type": FORMATS[fmt][0], "size": len(payload), "created_at": datetime.now(timezone.utc).isoformat(),
                    "filename": filename, "path": str(path), "preview_path": preview_path}
        meta_path = directory / f".{artifact_id}.meta.json"
        meta_path.write_text(json.dumps(metadata, ensure_ascii=False), encoding="utf-8")
        try:
            meta_path.chmod(0o600)
            path.chmod(0o600)
        except OSError:
            pass
        with self._lock:
            self._items[artifact_id] = metadata
        return self.public(metadata)

    @staticmethod
    def public(item: dict[str, Any]) -> dict[str, Any]:
        return {k: v for k, v in item.items() if k not in {"path", "preview_path", "workspace_id"}} | {
            "preview_url": f"/v1/artifacts/{item['id']}/preview",
            "download_url": f"/v1/artifacts/{item['id']}/download",
        }

    def get(self, workspace_id: str, artifact_id: str) -> dict[str, Any]:
        with self._lock:
            item = self._items.get(artifact_id)
            if not item or item["workspace_id"] != workspace_id:
                raise NotFound("artifact not found", details={"artifact_id": artifact_id})
            return dict(item)

    def list(self, workspace_id: str, limit: int = 50) -> list[dict[str, Any]]:
        with self._lock:
            rows = [dict(x) for x in self._items.values() if x["workspace_id"] == workspace_id]
        rows.sort(key=lambda x: x["created_at"], reverse=True)
        return [self.public(x) for x in rows[:limit]]
